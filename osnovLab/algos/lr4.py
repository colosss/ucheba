import random
from typing import List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from hash_func import hash_index

def _rand_key() -> str:
    letters = 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'
    digits = '0123456789'
    return (
        random.choice(letters) +
        random.choice(letters) +
        random.choice(digits) +
        random.choice(digits) +
        random.choice(letters) +
        random.choice(letters)
    )

def run_experiment(segments: int, multiplier: float = 2.5, seed: Optional[int] = None
                  ) -> Tuple[List[int], int]:
    """
    Формирует N = round(segments * multiplier) случайных ключей,
    считает, сколько раз адресовался каждый сегмент.
    Возвращает (counts, N).
    """
    if seed is not None:
        random.seed(seed)
    n = max(1, int(round(segments * float(multiplier))))
    counts = [0] * segments
    for _ in range(n):
        idx = hash_index(_rand_key(), segments)
        counts[idx] += 1
    return counts, n

def save_bucket_barchart_excel(counts: List[int], path_xlsx: str, title: str = None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = 'data'
    ws.append(['bucket', 'count'])
    for i, c in enumerate(counts):
        ws.append([i, c])

    chart = BarChart()
    chart.type = 'col'
    chart.title = title or f'Counts per bucket (m={len(counts)})'
    chart.x_axis.title = 'Bucket index'
    chart.y_axis.title = 'Count'

    min_row = 2
    max_row = ws.max_row
    cats = Reference(ws, min_col=1, min_row=min_row, max_row=max_row)
    data = Reference(ws, min_col=2, min_row=min_row, max_row=max_row)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)
    chart.legend = None

    ws.add_chart(chart, 'E2')
    wb.save(path_xlsx)
